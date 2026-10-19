import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SCREENSHOT_DIR = os.path.abspath("docs/screenshots")
DOCX_OUT = os.path.abspath("docs/ErgoSnap_Analytics_Project_Report.docx")
HTML_OUT = os.path.abspath("docs/ErgoSnap_Analytics_Project_Report.html")
PDF_OUT = os.path.abspath("docs/ErgoSnap_Analytics_Project_Report.pdf")
CHECKLIST_OUT = os.path.abspath("docs/ErgoSnap_Project_Evidence_Checklist.md")
ARTIFACT_MD = r"C:\Users\ruthi\.gemini\antigravity\brain\e70f7897-e8d8-4cc0-bfe2-25a9d6950c67\report.md"
ARTIFACT_IMG_DIR = r"/C:/Users/ruthi/.gemini/antigravity/brain/e70f7897-e8d8-4cc0-bfe2-25a9d6950c67/screenshots"

def style_table(table, header_bg="1E3A8A", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
        if r_idx == 0:
            trPr.append(parse_xml(r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'''
                <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)

            if r_idx == 0:
                shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{header_bg}"/>')
                tcPr.append(shd)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9)
            elif r_idx % 2 == 1:
                shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{alt_bg}"/>')
                tcPr.append(shd)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
            else:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)

print("Deliverable helper loaded successfully.")
