import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SCREENSHOT_DIR = os.path.abspath("docs/screenshots")
DOCX_OUT = os.path.abspath("docs/ErgoSnap_Project_Report.docx")
MD_OUT = os.path.abspath("docs/ErgoSnap_Project_Report.md")
ARTIFACT_MD = r"C:\Users\ruthi\.gemini\antigravity\brain\e70f7897-e8d8-4cc0-bfe2-25a9d6950c67\report.md"

def build_docx(markdown_content, docx_path):
    print("Compiling Word Document (.docx)...")
    doc = docx.Document()
    
    # Page setup: Standard Letter/A4, 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ErgoSnap Analytics — Academic Project Report | Production Engineering")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Department of Production Engineering — B.E. Capstone Project Report")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Arial'
    normal_font.size = Pt(10)
    normal_font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    lines = markdown_content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Code block handling
        if line.strip().startswith('```'):
            if in_code_block:
                # Flush code block
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.right_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                
                # Add shaded border box via XML
                pPr = p._p.get_or_add_pPr()
                pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:left w:val="single" w:sz="18" w:space="8" w:color="DC2626"/></w:pBdr>')
                pPr.append(pBdr)
                shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F1F5F9"/>')
                pPr.append(shd)

                code_text = '\n'.join(code_lines)
                c_run = p.add_run(code_text)
                c_run.font.name = 'Consolas'
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = RGBColor(15, 23, 42)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Image tag: ![Caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1)
            img_rel = img_match.group(2)
            img_filename = os.path.basename(img_rel)
            img_path = os.path.join(SCREENSHOT_DIR, img_filename)

            if os.path.exists(img_path):
                # Add image paragraph
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(3)
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Inches(5.8))

                # Add caption paragraph
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(12)
                r_cap = p_cap.add_run(f"{caption}")
                r_cap.font.name = "Arial"
                r_cap.font.size = Pt(8.5)
                r_cap.font.italic = True
                r_cap.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Major Document Title
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title_text)
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Midnight Slate
            i += 1
            continue

        # Section Heading 1 (e.g. Page 1 — Introduction...)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.font.name = 'Arial'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(220, 38, 38) # Red Accent
            i += 1
            continue

        # Sub-heading 2
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h_text)
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Horizontal Divider
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            pPr.append(pBdr)
            i += 1
            continue

        # Bullet List Item
        if stripped.startswith('- ') or stripped.startswith('* '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            
            # Simple bold parsing: **text**
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Numbered List Item
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            num_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Regular Paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('*') and part.endswith('*') and len(part) > 2:
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                else:
                    p.add_run(part)

        i += 1

    doc.save(docx_path)
    print(f"[OK] Word Document saved to: {docx_path}")

print("Build script helper loaded successfully.")
